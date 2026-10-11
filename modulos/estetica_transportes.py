# modulos/estetica_transportes.py
"""
Módulo Estética y Transportes

Ventana con buscador de clientes, ficha de transporte y ficha de estética
(mascotas). Importa Excel la primera vez y guarda todo en JSON.
"""

import os
import json
import re
import unicodedata
import math
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox

import ttkbootstrap as ttk
from ttkbootstrap.constants import *

from openpyxl import load_workbook

from core.rutas_estetica import (
    CARPETA_DATOS,
    CLIENTES_MAESTRO,
    CLIENTES_EDICIONES,
    MASCOTAS_FILE,
    SERVICIOS_ESTETICA,
    TARIFAS_TRANSPORTE,
    TRANSPORTES_CLIENTE,
)
from core.geocoding import geocodificar
from ui.utils import configurar_ventana
from config.veterinaria import sucursal_activa as _sucursal_activa


# ---------------------------------------------------------------
# Normalización
# ---------------------------------------------------------------

ACRONIMOS = {
    "mvz", "sa", "sc", "cv", "sapi", "rl", "ac", "sp", "dr", "dra",
    "de", "y", "la", "el", "los", "las",
}


def quitar_acentos(texto: str) -> str:
    if not texto:
        return ""
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def normalizar_nombre(nombre: str) -> str:
    if not nombre:
        return ""
    palabras = nombre.strip().split()
    resultado = []
    for p in palabras:
        limpio = p.strip(".,").lower()
        if limpio in ACRONIMOS:
            resultado.append(p.upper() if len(limpio) <= 3 else p.capitalize())
        else:
            resultado.append(p.capitalize())
    return " ".join(resultado)


def clave_busqueda(texto: str) -> str:
    return quitar_acentos((texto or "").lower())


# ---------------------------------------------------------------
# JSON helpers
# ---------------------------------------------------------------

def _cargar_json(path: Path, default):
    if not path.exists():
        return default
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def _guardar_json(path: Path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------
# Traducción de códigos de estética
# ---------------------------------------------------------------

_PATRONES_ESTETICA = [
    (r"^B\s+GDE\s+PL.*", "Baño grande pelo largo"),
    (r"^B\s+GDE\s+PM.*", "Baño grande pelo mediano"),
    (r"^B\s+GDE\s+PC.*", "Baño grande pelo corto"),
    (r"^B\s+GIG\s+PL.*", "Baño gigante pelo largo"),
    (r"^B\s+GIG\s+PC.*", "Baño gigante pelo corto"),
    (r"^B\s+GATO\s+G\s+PL.*", "Baño gato grande pelo largo"),
    (r"^B\s+GATO\s+G\s+PC.*", "Baño gato grande pelo corto"),
    (r"^B\s+GATO\s+CH\s+PL.*", "Baño gato chico pelo largo"),
    (r"^B\s+GATO\s+CH\s+PC.*", "Baño gato chico pelo corto"),
    (r"^B\s+M\s+PM.*", "Baño mediano pelo mediano"),
    (r"^B\s+M\s+PL.*", "Baño mediano pelo largo"),
    (r"^B\s+M\s+PC.*", "Baño mediano pelo corto"),
    (r"^B\s+CH\s+PL.*", "Baño chico pelo largo"),
    (r"^B\s+CH\s+PC.*", "Baño chico pelo corto"),
    (r"^B\s+EX\s+CH\s+PC.*", "Baño extra chico pelo corto"),

    (r"^BG\s+A\s+GIG\s+PC.*", "Baño garrapaticida A gigante pelo corto"),
    (r"^BG\s+A\s+G\s+PC.*",   "Baño garrapaticida A grande pelo corto"),
    (r"^BG\s+A\s+M\s+PC.*",   "Baño garrapaticida A mediano pelo corto"),
    (r"^BG\s+A\s+CH\s+PC.*",  "Baño garrapaticida A chico pelo corto"),
    (r"^BG\s+B\s+GIG\s+PC.*", "Baño garrapaticida B gigante pelo corto"),
    (r"^BG\s+B\s+G\s+PC.*",   "Baño garrapaticida B grande pelo corto"),
    (r"^BG\s+B\s+M\s+PC.*",   "Baño garrapaticida B mediano pelo corto"),
    (r"^BG\s+B\s+CH\s+PC.*",  "Baño garrapaticida B chico pelo corto"),
    (r"^BG\s+C\s+GIG\s+PC.*", "Baño garrapaticida C gigante pelo corto"),
    (r"^BG\s+C\s+G\s+PC.*",   "Baño garrapaticida C grande pelo corto"),
    (r"^BG\s+C\s+M\s+PC.*",   "Baño garrapaticida C mediano pelo corto"),
    (r"^BG\s+C\s+CH\s+PC.*",  "Baño garrapaticida C chico pelo corto"),
    (r"^BG\s+D\s+GIG\s+PC.*", "Baño garrapaticida D gigante pelo corto"),
    (r"^BG\s+D\s+G\s+PC.*",   "Baño garrapaticida D grande pelo corto"),
    (r"^BG\s+D\s+M\s+PC.*",   "Baño garrapaticida D mediano pelo corto"),
    (r"^BG\s+D\s+CH\s+PC.*",  "Baño garrapaticida D chico pelo corto"),

    (r"^EX\s+G\s+A\s+GIG\s+PL.*", "Extra garrapaticida A gigante pelo largo"),
    (r"^EX\s+G\s+A\s+G\s+PL.*",   "Extra garrapaticida A grande pelo largo"),
    (r"^EX\s+G\s+A\s+G\s+PM.*",   "Extra garrapaticida A grande pelo mediano"),
    (r"^EX\s+G\s+A\s+M\s+PL.*",   "Extra garrapaticida A mediano pelo largo"),
    (r"^EX\s+G\s+A\s+M\s+PM.*",   "Extra garrapaticida A mediano pelo mediano"),
    (r"^EX\s+G\s+A\s+CH\s+PL.*",  "Extra garrapaticida A chico pelo largo"),

    (r"^CP\s+GIG$",      "Corte pelo gigante"),
    (r"^CP\s+G\s+PM$",   "Corte pelo grande pelo mediano"),
    (r"^CP\s+G$",        "Corte pelo grande"),
    (r"^CP\s+M\s+PM$",   "Corte pelo mediano pelo mediano"),
    (r"^CP\s+M$",        "Corte pelo mediano"),
    (r"^CP\s+CH$",       "Corte pelo chico"),
    (r"^CP\s+GATO\s+G$", "Corte pelo gato grande"),
    (r"^CP\s+GATO\s+CH$","Corte pelo gato chico"),

    (r"^CORTE\s+DE\s+UÑAS\s+A$", "Corte de uñas A"),
    (r"^CORTE\s+DE\s+UÑAS\s+B$", "Corte de uñas B"),
    (r"^CORTE\s+DE\s+UÑAS\s+C$", "Corte de uñas C"),
    (r"^QUICK\s+STOP.*", "Corte de uñas - Quick Stop"),

    (r"^EX\s+BAÑO\s+MEDICADO.*", "Baño medicado"),
    (r"^EX\s+NUDOS\s+O\s+DESLANADOS\s+(\d)", r"Deslanado nivel \1"),
    (r"^EX\s+NUDOS\s+O\s+DESLANADOS\s+SESION", "Deslanado sesión"),
    (r"^EX\s+SHAMPOO\s+(.+)", r"Shampoo \1"),
    (r"^SERVICIO\s+EXTRA\s+ECTO\s+(.+)", r"Fumigación \1"),
    (r"^SERVICIO\s+EXTRA\s+([A-E])$", r"Servicio extra \1"),
    (r"^PEINADO\s+(.+)", r"Peinado \1"),
    (r"^TINTE\s+PERMANENTE", "Tinte permanente"),
    (r"^TRATAMIENTO\s+(.+)", r"Tratamiento \1"),
    (r"^TALCO\s+OTICO\s+(.+)", r"Talco ótico \1"),
]


def traducir_codigo_estetica(codigo: str) -> str:
    codigo_norm = codigo.strip().upper()
    for patron, trad in _PATRONES_ESTETICA:
        if re.match(patron, codigo_norm):
            return trad
    return codigo


def duracion_por_defecto(familia: str, codigo: str) -> int:
    fam = familia.upper()
    if fam == "CORTE DE PELO":
        return 60
    if fam == "BAÑOS":
        return 45
    if fam == "CORTE DE UÑAS":
        return 15
    if fam == "SERVICIOS EXTRAS":
        return 20
    return 30


# ---------------------------------------------------------------
# Haversine
# ---------------------------------------------------------------

def haversine_km(lat1, lon1, lat2, lon2):
    R = 6371.0
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)
    a = (math.sin(dlat / 2) ** 2
         + math.cos(math.radians(lat1)) * math.cos(math.radians(lat2))
         * math.sin(dlon / 2) ** 2)
    return R * 2 * math.asin(math.sqrt(a))


# ---------------------------------------------------------------
# Cálculo de ruta completa (ida + vuelta + recolección)
# ---------------------------------------------------------------

# Ajustar con pruebas reales (ver README del módulo)
FACTOR_RUTA = 1.4
VELOCIDAD_KMH = 22.0
MINUTOS_RECOLECCION = 5


def calcular_ruta_completa(lat_cliente, lng_cliente):
    """
    Calcula la ruta desde la sucursal activa hasta el cliente.
    """
    suc = _sucursal_activa()
    coords = suc.get("coordenadas", {})
    lat = coords.get("lat")
    lng = coords.get("lng")

    if None in (lat_cliente, lng_cliente, lat, lng):
        return None, None, None

    distancia_recta = haversine_km(lat, lng, lat_cliente, lng_cliente)
    distancia_km = round(distancia_recta * FACTOR_RUTA, 2)

    tiempo_ida_min = max(1, round((distancia_km / VELOCIDAD_KMH) * 60))
    tiempo_total_min = tiempo_ida_min * 2 + MINUTOS_RECOLECCION

    return distancia_km, tiempo_ida_min, tiempo_total_min


# ---------------------------------------------------------------
# Detección de tipo de Excel por contenido
# ---------------------------------------------------------------

COLS_CLIENTES_MASCOTAS = {
    "CLIENTE", "MASCOTA", "ESPECIE", "RAZA",
    "COMUNICACION_IDCLIENTE", "COMUNICACION_IDANIMAL",
}


def _detectar_clientes_mascotas(encabezados: dict) -> bool:
    """True si el Excel trae cliente + mascota en la misma hoja."""
    return COLS_CLIENTES_MASCOTAS.issubset(encabezados.keys())

def _leer_encabezados(ws):
    """Devuelve dict {NOMBRE_COLUMNA: índice} normalizado."""
    encabezados = {}
    for idx, cell in enumerate(ws[1]):
        if cell.value:
            key = quitar_acentos(str(cell.value).strip().upper())
            encabezados[key] = idx
    return encabezados


def _detectar_tipo_excel(ruta_excel: str):
    """
    Detecta el tipo de Excel por sus columnas.
    Devuelve: 'clientes', 'servicios_estetica', 'tarifas_transporte' o None
    """
    try:
        wb = load_workbook(ruta_excel, data_only=True, read_only=True)
        ws = wb.active
        encabezados = _leer_encabezados(ws)

        # Clientes + mascotas (una fila por mascota)
        if _detectar_clientes_mascotas(encabezados):
            wb.close()
            return "clientes"

        # Precios (servicios/tarifas): DESCRIPCION + SECCION + FAMILIA + PVP
        if all(c in encabezados for c in ("DESCRIPCION", "SECCION", "FAMILIA", "PVP")):
            idx_sec = encabezados["SECCION"]
            secciones = set()
            for fila in ws.iter_rows(min_row=2, values_only=True):
                if fila and len(fila) > idx_sec and fila[idx_sec]:
                    secciones.add(str(fila[idx_sec]).strip().upper())
            wb.close()
            if "ESTETICA" in secciones or "ESTÉTICA" in secciones:
                return "servicios_estetica"
            if "TRANSPORTE" in secciones:
                return "tarifas_transporte"

        wb.close()
    except Exception as e:
        print(f"[detectar_tipo] Error: {e}")
    return None


def _preview_excel(ruta_excel: str) -> dict:
    """Info resumida para mostrar antes de importar."""
    try:
        wb = load_workbook(ruta_excel, data_only=True, read_only=True)
        ws = wb.active
        total = ws.max_row - 1 if ws.max_row else 0
        wb.close()
    except Exception:
        total = "?"
    return {"tipo": _detectar_tipo_excel(ruta_excel), "total_filas": total}


# ---------------------------------------------------------------
# Importadores con merge (solo añaden nuevos)
# ---------------------------------------------------------------

def importar_clientes_mascotas_excel(ruta_excel: str) -> dict:
    """
    Importa Excel con una fila por mascota.

    - Cliente: identifica por COMUNICACION_IDCLIENTE.
      * Si ya existe, fusiona datos faltantes (id_cliente, telefono, cp).
      * Si no existe, lo crea.
    - Mascota: identifica por COMUNICACION_IDANIMAL dentro del cliente.
      * Si ya existe, se salta.
      * Si no existe, se agrega (aunque se llame igual que otra).
    """
    wb = load_workbook(ruta_excel, data_only=True)
    ws = wb.active
    enc = _leer_encabezados(ws)

    idx_cli   = enc["CLIENTE"]
    idx_mas   = enc["MASCOTA"]
    idx_dir   = enc.get("DIRECCIÓN", enc.get("DIRECCION"))
    idx_pob   = enc.get("POBLACIÓN", enc.get("POBLACION"))
    idx_cp    = enc.get("CP")
    idx_tel   = enc.get("TELÉFONO1", enc.get("TELEFONO1"))
    idx_esp   = enc.get("ESPECIE")
    idx_raza  = enc.get("RAZA")
    idx_idcli = enc["COMUNICACION_IDCLIENTE"]
    idx_idani = enc["COMUNICACION_IDANIMAL"]

    maestro  = _cargar_json(CLIENTES_MAESTRO, [])
    mascotas = _cargar_json(MASCOTAS_FILE, {})

    # Índice por id_cliente (solo los que ya lo tengan)
    por_id_cli = {
        str(c.get("id_cliente")).strip(): c
        for c in maestro if c.get("id_cliente")
    }
    max_id = max((c.get("id", 0) for c in maestro), default=0)

    # Índice por nombre normalizado (para fusionar con clientes viejos sin id_cliente)
    por_nombre = {
        clave_busqueda(c.get("cliente", "")): c
        for c in maestro
    }

    stats = {
        "clientes_nuevos": 0,
        "clientes_fusionados": 0,
        "mascotas_nuevas": 0,
        "mascotas_existentes": 0,
    }

    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila or len(fila) <= max(idx_idcli, idx_idani, idx_cli, idx_mas):
            continue

        id_cli = str(fila[idx_idcli] or "").strip()
        nombre_raw = str(fila[idx_cli] or "").strip()
        if not id_cli or not nombre_raw:
            continue

        nombre_norm = normalizar_nombre(nombre_raw)
        clave_norm  = clave_busqueda(nombre_raw)

        # ---- Cliente: buscar por ID, luego por nombre ----
        cliente = por_id_cli.get(id_cli) or por_nombre.get(clave_norm)

        if cliente is None:
            # Cliente nuevo
            max_id += 1
            cliente = {
                "id": max_id,
                "id_cliente": id_cli,
                "cliente": nombre_norm,
                "cliente_raw": nombre_raw,
                "poblacion": "",
                "direccion": "",
                "cp": "",
                "telefono": "",
                "notas_cliente": "",
                "duplicado": False,
            }
            maestro.append(cliente)
            stats["clientes_nuevos"] += 1
        else:
            # Fusionar datos faltantes
            if not cliente.get("id_cliente"):
                cliente["id_cliente"] = id_cli
            stats["clientes_fusionados"] += 1

        # Registrar en índices
        por_id_cli[id_cli] = cliente
        por_nombre[clave_norm] = cliente

        # ---- Completar campos que falten ----
        dir_val = str(fila[idx_dir] or "").strip() if idx_dir is not None else ""
        cp_val  = str(fila[idx_cp]  or "").strip() if idx_cp  is not None else ""
        tel_val = str(fila[idx_tel] or "").strip() if idx_tel is not None else ""
        pob_val = str(fila[idx_pob] or "").strip() if idx_pob is not None else ""

        if dir_val and not cliente.get("direccion"):
            cliente["direccion"] = dir_val
        if cp_val and not cliente.get("cp"):
            cliente["cp"] = cp_val
        if tel_val and not cliente.get("telefono"):
            cliente["telefono"] = tel_val
        if pob_val and not cliente.get("poblacion"):
            cliente["poblacion"] = pob_val
        # ---- Mascota ----
        nombre_cli = cliente["cliente"]
        lista = mascotas.setdefault(nombre_cli, [])

        id_ani = str(fila[idx_idani] or "").strip()
        nombre_mas = str(fila[idx_mas] or "").strip() or "[sin nombre]"

        # Dedupe por id_animal (si ya existe ese ID en la lista, skip)
        if id_ani and any(
            str(m.get("id_animal", "")).strip() == id_ani for m in lista
        ):
            stats["mascotas_existentes"] += 1
            continue

        lista.append({
            "nombre": nombre_mas,
            "especie": str(fila[idx_esp] or "").strip() if idx_esp is not None else "",
            "raza":    str(fila[idx_raza] or "").strip() if idx_raza is not None else "",
            "id_animal": id_ani,
            "caracter": "",
            "precauciones": "",
            "notas": "",
        })
        stats["mascotas_nuevas"] += 1

    wb.close()
    _guardar_json(CLIENTES_MAESTRO, maestro)
    _guardar_json(MASCOTAS_FILE, mascotas)

    stats["total_maestro"] = len(maestro)
    stats["total_mascotas"] = sum(len(v) for v in mascotas.values())
    return stats


def importar_servicios_estetica(ruta_excel: str) -> dict:
    """Importa servicios de estética. Solo añade los que no existan."""
    wb = load_workbook(ruta_excel, data_only=True)
    ws = wb.active
    encabezados = _leer_encabezados(ws)

    idx_desc = encabezados.get("DESCRIPCION", 0)
    idx_fam = encabezados.get("FAMILIA", 2)
    idx_sub = encabezados.get("SUBFAMILIA", 3)
    idx_pvp = encabezados.get("PVP", 4)

    FAMILIAS_SERVICIO = {"BAÑOS", "CORTE DE PELO", "CORTE DE UÑAS", "SERVICIOS EXTRAS"}

    existentes_json = _cargar_json(SERVICIOS_ESTETICA, [])
    codigos_existentes = {s.get("codigo", "").strip().upper() for s in existentes_json}
    max_id = max((s.get("id", 0) for s in existentes_json), default=0)

    nuevos = 0
    existentes = 0

    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila or len(fila) <= idx_desc or not fila[idx_desc]:
            continue
        familia = str(fila[idx_fam] or "").strip().upper() if len(fila) > idx_fam else ""
        if familia not in FAMILIAS_SERVICIO:
            continue
        codigo = str(fila[idx_desc]).strip()
        if codigo.upper() in codigos_existentes:
            existentes += 1
            continue
        try:
            pvp = float(fila[idx_pvp] or 0) if len(fila) > idx_pvp else 0.0
        except (TypeError, ValueError):
            pvp = 0.0
        max_id += 1
        existentes_json.append({
            "id": max_id,
            "codigo": codigo,
            "descripcion": traducir_codigo_estetica(codigo),
            "familia": familia,
            "subfamilia": str(fila[idx_sub] or "").strip() if len(fila) > idx_sub else "",
            "precio": pvp,
            "duracion_min": duracion_por_defecto(familia, codigo),
            "activo": True,
        })
        codigos_existentes.add(codigo.upper())
        nuevos += 1

    _guardar_json(SERVICIOS_ESTETICA, existentes_json)
    return {
        "nuevos": nuevos,
        "existentes": existentes,
        "total_maestro": len(existentes_json),
    }


def importar_tarifas_transporte(ruta_excel: str) -> dict:
    """Importa tarifas. Solo añade las que no existan."""
    wb = load_workbook(ruta_excel, data_only=True)
    ws = wb.active
    encabezados = _leer_encabezados(ws)

    idx_desc = encabezados.get("DESCRIPCION", 0)
    idx_fam = encabezados.get("FAMILIA", 2)
    idx_sub = encabezados.get("SUBFAMILIA", 3)
    idx_pvp = encabezados.get("PVP", 4)

    existentes_json = _cargar_json(TARIFAS_TRANSPORTE, [])
    codigos_existentes = {t.get("codigo", "").strip().upper() for t in existentes_json}
    max_id = max((t.get("id", 0) for t in existentes_json), default=0)

    nuevos = 0
    existentes = 0

    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila or len(fila) <= idx_desc or not fila[idx_desc]:
            continue
        codigo = str(fila[idx_desc]).strip()
        if codigo.upper() in codigos_existentes:
            existentes += 1
            continue
        try:
            precio = float(fila[idx_pvp] or 0) if len(fila) > idx_pvp else 0.0
        except (TypeError, ValueError):
            precio = 0.0
        max_id += 1
        existentes_json.append({
            "id": max_id,
            "codigo": codigo,
            "familia": str(fila[idx_fam] or "").strip() if len(fila) > idx_fam else "",
            "subfamilia": str(fila[idx_sub] or "").strip() if len(fila) > idx_sub else "",
            "precio": precio,
            "activo": True,
        })
        codigos_existentes.add(codigo.upper())
        nuevos += 1

    _guardar_json(TARIFAS_TRANSPORTE, existentes_json)
    return {
        "nuevos": nuevos,
        "existentes": existentes,
        "total_maestro": len(existentes_json),
    }


# ---------------------------------------------------------------
# Ventana principal del módulo
# ---------------------------------------------------------------

class AppEsteticaTransportes(ttk.Toplevel):
    """Ventana principal del módulo Estética y Transportes."""

    def __init__(self, master=None):
        super().__init__(master)
        self.title("Vet Suite — Estética y Transportes")

        configurar_ventana(
            master, self,
            ancho=1300, alto=880,
            min_ancho=1050, min_alto=700,
            centrar_en_padre=True,
        )

        self.clientes = _cargar_json(CLIENTES_MAESTRO, [])
        self.mascotas = _cargar_json(MASCOTAS_FILE, {})
        self.ediciones = _cargar_json(CLIENTES_EDICIONES, {})
        self.transportes = _cargar_json(TRANSPORTES_CLIENTE, {})

        self._construir_ui()
        self._refrescar_tabla()

    def _construir_ui(self):
        # ---- Barra superior con buscador ----
        top = ttk.LabelFrame(self, text="Búsqueda", padding=10)
        top.pack(fill="x", padx=10, pady=5)

        ttk.Label(top, text="🔍 Buscar:").grid(row=0, column=0, padx=5, sticky="e")
        self.var_busqueda = tk.StringVar()
        self.var_busqueda.trace_add("write", lambda *_: self._refrescar_tabla())
        entry = ttk.Entry(top, textvariable=self.var_busqueda, width=50,
                          style="Custom.TEntry")
        entry.grid(row=0, column=1, padx=5)
        entry.focus_set()

        ttk.Label(top,
                  text="(busca en nombre, población, dirección y mascotas)",
                  foreground="gray").grid(row=0, column=2, padx=10, sticky="w")

        # ---- Barra de botones ----
        botones = ttk.Frame(self, padding=(10, 5))
        botones.pack(fill="x", padx=10)

        self._btn_importar = ttk.Button(
            botones, text="📥 Importar Excel",
            command=self._importar_exceles,
            bootstyle="info-outline",
        )
        self._btn_importar.pack(side="left", padx=3)

        ttk.Button(
            botones, text="📚 Catálogos",
            command=self._abrir_catalogos,
            bootstyle="secondary-outline",
        ).pack(side="left", padx=3)

        ttk.Button(
            botones, text="⚠️ Duplicados / Basura",
            command=self._abrir_duplicados_basura,
            bootstyle="warning-outline",
        ).pack(side="left", padx=3)

        ttk.Button(
            botones, text="+ Nuevo cliente",
            command=self._nuevo_cliente,
            bootstyle="success-outline",
        ).pack(side="left", padx=3)

        self._actualizar_boton_importar()

        # ---- Tabla ----
        frame_tabla = ttk.LabelFrame(self, text="Clientes", padding=5)
        frame_tabla.pack(fill="both", expand=True, padx=10, pady=5)

        contenedor = ttk.Frame(frame_tabla)
        contenedor.pack(fill="both", expand=True)

        cols = ("alerta", "cliente", "poblacion", "telefono", "direccion", "mascotas", "transporte")
        encabezados = {
            "alerta": "⚠️",
            "cliente": "CLIENTE",
            "poblacion": "POBLACIÓN",
            "telefono": "TELÉFONO",
            "direccion": "DIRECCIÓN",
            "mascotas": "MASCOTAS",
            "transporte": "TRANSPORTE",
        }
        anchos = {
            "alerta": 40,
            "cliente": 230,
            "poblacion": 150,
            "telefono": 110,
            "direccion": 300,
            "mascotas": 200,
            "transporte": 130,
        }

        self.tabla = ttk.Treeview(contenedor, columns=cols, show="headings",
                                  height=20, bootstyle="primary")
        for c in cols:
            self.tabla.heading(c, text=encabezados[c])
            self.tabla.column(c, width=anchos[c],
                              anchor="center" if c in ("alerta", "transporte") else "w")

        sb_v = ttk.Scrollbar(contenedor, orient="vertical", command=self.tabla.yview)
        sb_h = ttk.Scrollbar(contenedor, orient="horizontal", command=self.tabla.xview)
        self.tabla.configure(yscrollcommand=sb_v.set, xscrollcommand=sb_h.set)

        self.tabla.grid(row=0, column=0, sticky="nsew")
        sb_v.grid(row=0, column=1, sticky="ns")
        sb_h.grid(row=1, column=0, sticky="ew")
        contenedor.rowconfigure(0, weight=1)
        contenedor.columnconfigure(0, weight=1)

        self.tabla.bind("<Double-Button-1>", self._abrir_detalle)

        self.lbl_info = ttk.Label(self, text="", font=("Segoe UI", 10, "bold"))
        self.lbl_info.pack(fill="x", padx=15, pady=(0, 8))

    # ---------- Botón importar/reimportar ----------

    def _actualizar_boton_importar(self):
        """Cambia el texto del botón según si ya hay registros."""
        try:
            if CLIENTES_MAESTRO.exists():
                self._btn_importar.configure(text="📥 Reimportar Excel")
            else:
                self._btn_importar.configure(text="📥 Importar Excel")
        except Exception:
            pass

    # ---------- Tabla ----------

    def _refrescar_tabla(self):
        """Recarga los JSONs desde disco y repinta la tabla."""
        self.clientes = _cargar_json(CLIENTES_MAESTRO, [])
        self.mascotas = _cargar_json(MASCOTAS_FILE, {})
        self.transportes = _cargar_json(TRANSPORTES_CLIENTE, {})

        for i in self.tabla.get_children():
            self.tabla.delete(i)

        query = clave_busqueda(self.var_busqueda.get())
        contador = 0

        for c in self.clientes:
            nombre = c.get("cliente", "")
            mascotas_cli = self.mascotas.get(nombre, [])
            blob = " ".join([
                nombre,
                c.get("poblacion", ""),
                c.get("direccion", ""),
                " ".join(m.get("nombre", "") for m in mascotas_cli),
            ])
            if query and query not in clave_busqueda(blob):
                continue

            # ¿Tiene notas en algún lado? → ⚠️
            tiene_notas = False
            if c.get("notas_cliente", "").strip():
                tiene_notas = True
            if self.mascotas.get(f"__notas__{nombre}", "").strip():
                tiene_notas = True
            trans = self.transportes.get(nombre, {})
            if trans.get("notas", "").strip():
                tiene_notas = True
            if not tiene_notas:
                for m in mascotas_cli:
                    if m.get("notas", "").strip():
                        tiene_notas = True
                        break

            alerta = "⚠️" if tiene_notas else ""

            mascotas_txt = ", ".join(m.get("nombre", "") for m in mascotas_cli) or "—"
            if trans:
                dist = trans.get("distancia_km", "")
                # Preferimos tiempo_total_min si existe, si no tiempo_min
                tiempo_total = trans.get("tiempo_total_min",
                                          trans.get("tiempo_min", ""))
                trans_txt = f"{dist} km / {tiempo_total} min" if dist else "sin datos"
            else:
                trans_txt = "sin datos"

            self.tabla.insert("", END, iid=str(c.get("id")), values=(
                alerta,
                nombre,
                c.get("poblacion", ""),
                c.get("direccion", ""),
                c.get("telefono", ""),
                mascotas_txt,
                trans_txt,
            ))
            contador += 1

        self.lbl_info.config(text=f"{contador} clientes encontrados")

    # ---------- Importar ----------

    def _importar_exceles(self):
        """
        Importador inteligente:
        - Pide 1 o más Excel
        - Detecta tipo por contenido
        - Muestra preview
        - Importa con merge (solo añade nuevos)
        """
        rutas = filedialog.askopenfilenames(
            title="Selecciona los Excel a importar",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not rutas:
            return

        analisis = []
        for r in rutas:
            nombre = os.path.basename(r)
            info = _preview_excel(r)
            analisis.append({
                "ruta": r,
                "nombre": nombre,
                "tipo": info.get("tipo"),
                "total": info.get("total_filas", "?"),
            })

        # Preview
        etiquetas_tipo = {
            "clientes": "Clientes",
            "servicios_estetica": "Servicios de estética",
            "tarifas_transporte": "Tarifas de transporte",
        }

        lineas = ["Archivos detectados:\n"]
        for a in analisis:
            if a["tipo"]:
                etiqueta = etiquetas_tipo.get(a["tipo"], a["tipo"])
                lineas.append(f"  • {a['nombre']} → {etiqueta} ({a['total']} filas)")
            else:
                lineas.append(f"  • {a['nombre']} → ❌ NO RECONOCIDO")

        lineas.append("\nSe añadirán solo los registros que no existan.")
        lineas.append("¿Continuar?")

        if not messagebox.askyesno("Preview de importación", "\n".join(lineas),
                                   parent=self):
            return

        resumen = []
        for a in analisis:
            if not a["tipo"]:
                resumen.append(f"❌ {a['nombre']}: no reconocido")
                continue
            try:
                if a["tipo"] == "clientes":
                    r = importar_clientes_mascotas_excel(a["ruta"])
                    resumen.append(
                        f"✅ Clientes + mascotas:\n"
                        f"   • {r['clientes_nuevos']} clientes nuevos, "
                        f"{r['clientes_fusionados']} fusionados\n"
                        f"   • {r['mascotas_nuevas']} mascotas nuevas, "
                        f"{r['mascotas_existentes']} ya existían\n"
                        f"   • Total: {r['total_maestro']} clientes, "
                        f"{r['total_mascotas']} mascotas"
                    )
                elif a["tipo"] == "servicios_estetica":
                    r = importar_servicios_estetica(a["ruta"])
                    resumen.append(
                        f"✅ Servicios estética: +{r['nuevos']} nuevos "
                        f"({r['existentes']} ya existían)"
                    )
                elif a["tipo"] == "tarifas_transporte":
                    r = importar_tarifas_transporte(a["ruta"])
                    resumen.append(
                        f"✅ Tarifas transporte: +{r['nuevos']} nuevos "
                        f"({r['existentes']} ya existían)"
                    )
            except Exception as e:
                resumen.append(f"❌ {a['nombre']}: {e}")

        self._recargar_clientes()
        self.mascotas = _cargar_json(MASCOTAS_FILE, {})
        self.transportes = _cargar_json(TRANSPORTES_CLIENTE, {})

        messagebox.showinfo("Importación completa", "\n".join(resumen), parent=self)
        self._actualizar_boton_importar()

    # ---------- Botones secundarios ----------

    def _abrir_catalogos(self):
        from dialogos.catalogo_editor import DialogoCatalogo
        DialogoCatalogo(self)

    def _abrir_duplicados_basura(self):
        from dialogos.duplicados_basura import DialogoDuplicadosBasura
        DialogoDuplicadosBasura(self, self._recargar_clientes)

    def _recargar_clientes(self):
        self.clientes = _cargar_json(CLIENTES_MAESTRO, [])
        self._refrescar_tabla()

    # ---------- Detalle / Nuevo ----------

    def _abrir_detalle(self, event=None):
        sel = self.tabla.selection()
        if not sel:
            return
        cliente_id = int(sel[0])
        cliente = next((c for c in self.clientes if c.get("id") == cliente_id), None)
        if not cliente:
            return
        from dialogos.cliente_detalle import DialogoClienteDetalle
        DialogoClienteDetalle(self, cliente, on_guardar=self._refrescar_tabla)

    def _nuevo_cliente(self):
        from dialogos.cliente_detalle import DialogoClienteDetalle
        nuevo = {
            "id": max((c.get("id", 0) for c in self.clientes), default=0) + 1,
            "cliente": "",
            "cliente_raw": "",
            "poblacion": "",
            "direccion": "",
            "duplicado": False,
            "nuevo": True,
        }
        DialogoClienteDetalle(self, nuevo, on_guardar=self._recargar_clientes)