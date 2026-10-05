# -*- coding: utf-8 -*-
"""
excel/egresos.py
Genera los Excel de Egresos (PUE y PPD).
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# ESTILOS
# ============================================================
THIN = Side(border_style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
DERECHA = Alignment(horizontal="right", vertical="center")
IZQUIERDA = Alignment(horizontal="left", vertical="center")
BOLD_WHITE = Font(bold=True, color="FFFFFF", size=11)
BOLD_DARK = Font(bold=True, color="000000", size=10)
FUENTE_NORMAL = Font()
FUENTE_LINK = Font(color="0563C1", underline="single")
FORMATO_MONEDA = '"$"#,##0.00'
FORMATO_FECHA = "DD/MM/YYYY"

COLOR_PUE = "1F4E3D"
COLOR_PPD = "2E5A88"


# ============================================================
# GENERAR EXCEL PUE
# ============================================================
def generar_excel_pue(facturas, ruta_xlsx, mes_nombre, anio):
    """
    Genera el Excel de facturas PUE.
    
    Columnas:
      # | No. FACTURA | UUID | RFC EMISOR | NOMBRE O RAZÓN SOCIAL DEL EMISOR |
      FECHA | TOTAL | METODO DE PAGO | OBSERVACIONES
    """
    # Ordenar por número de línea
    facturas = sorted(facturas, key=lambda f: f.get("linea", 9999))
    wb = Workbook()
    ws = wb.active
    ws.title = f"PUE {mes_nombre.capitalize()} {anio}"[:31]

    # Encabezados
    encabezados = [
        ("A", "#", 6),
        ("B", "No. FACTURA", 20),
        ("C", "UUID", 38),
        ("D", "RFC EMISOR", 16),
        ("E", "NOMBRE O RAZÓN SOCIAL DEL EMISOR", 50),
        ("F", "FECHA", 12),
        ("G", "TOTAL", 14),
        ("H", "METODO DE PAGO", 16),
        ("I", "OBSERVACIONES", 30),
    ]

    for letra, titulo, ancho in encabezados:
        c = ws[f"{letra}1"]
        c.value = titulo
        c.font = BOLD_WHITE
        c.fill = PatternFill("solid", fgColor=COLOR_PUE)
        c.alignment = CENTRO
        c.border = BORDER
        ws.column_dimensions[letra].width = ancho

    # Filas
    fila = 2
    for f in facturas:
        # Usar el número de línea guardado, o el índice
        num_linea = f.get("linea", 0) or (fila - 1)
        ws[f"A{fila}"] = num_linea
        ws[f"A{fila}"].alignment = CENTRO
        ws[f"A{fila}"].border = BORDER

        # No. FACTURA (con enlace al PDF)
        c = ws[f"B{fila}"]
        c.value = f.get("folio", "")
        c.alignment = IZQUIERDA
        c.border = BORDER
        ruta_pdf = _obtener_ruta_pdf_egreso(f, num_linea)
        if ruta_pdf:
            c.hyperlink = ruta_pdf
            c.font = FUENTE_LINK

        # UUID
        ws[f"C{fila}"] = f.get("uuid", "")
        ws[f"C{fila}"].alignment = IZQUIERDA
        ws[f"C{fila}"].border = BORDER

        # RFC EMISOR
        ws[f"D{fila}"] = f.get("rfc_emisor", "")
        ws[f"D{fila}"].alignment = CENTRO
        ws[f"D{fila}"].border = BORDER

        # NOMBRE EMISOR
        ws[f"E{fila}"] = f.get("nombre_emisor", "")
        ws[f"E{fila}"].alignment = IZQUIERDA
        ws[f"E{fila}"].border = BORDER

        # FECHA
        ws[f"F{fila}"] = f.get("fecha", "")
        ws[f"F{fila}"].alignment = CENTRO
        ws[f"F{fila}"].border = BORDER

        # TOTAL
        ws[f"G{fila}"] = float(f.get("total", 0))
        ws[f"G{fila}"].number_format = FORMATO_MONEDA
        ws[f"G{fila}"].alignment = DERECHA
        ws[f"G{fila}"].border = BORDER

        # METODO DE PAGO
        ws[f"H{fila}"] = f.get("forma_pago_texto", "")
        ws[f"H{fila}"].alignment = CENTRO
        ws[f"H{fila}"].border = BORDER

        # OBSERVACIONES
        ws[f"I{fila}"] = f.get("observacion", "")
        ws[f"I{fila}"].alignment = IZQUIERDA
        ws[f"I{fila}"].border = BORDER

        fila += 1

    # Fila de totales
    ws[f"F{fila}"] = "TOTAL:"
    ws[f"F{fila}"].font = BOLD_DARK
    ws[f"F{fila}"].alignment = DERECHA
    ws[f"G{fila}"] = f"=SUM(G2:G{fila-1})"
    ws[f"G{fila}"].font = BOLD_DARK
    ws[f"G{fila}"].number_format = FORMATO_MONEDA
    ws[f"G{fila}"].alignment = DERECHA

    wb.save(ruta_xlsx)
    return ruta_xlsx


# ============================================================
# GENERAR EXCEL PPD
# ============================================================
def generar_excel_ppd(facturas, ruta_xlsx, mes_nombre, anio):
    """
    Genera el Excel de facturas PPD.
    
    Columnas:
      # | FECHA | FACTURA | QVET | PROVEEDOR | UUID | METODO |
      SUBTOTAL | IVA | IEPS | TOTAL | REFERENCIA
    """
    # Ordenar por número de línea (igual que PUE)
    facturas = sorted(facturas, key=lambda f: f.get("linea", 9999))

    wb = Workbook()
    ws = wb.active
    ws.title = f"PPD {mes_nombre.capitalize()} {anio}"[:31]

    encabezados = [
        ("A", "#", 6),
        ("B", "FECHA", 12),
        ("C", "FACTURA", 20),
        ("D", "QVET", 12),
        ("E", "PROVEEDOR", 50),
        ("F", "UUID", 38),
        ("G", "METODO", 10),
        ("H", "SUBTOTAL", 14),
        ("I", "IVA", 12),
        ("J", "IEPS", 12),
        ("K", "TOTAL", 14),
        ("L", "REFERENCIA", 20),
    ]

    for letra, titulo, ancho in encabezados:
        c = ws[f"{letra}1"]
        c.value = titulo
        c.font = BOLD_WHITE
        c.fill = PatternFill("solid", fgColor=COLOR_PPD)
        c.alignment = CENTRO
        c.border = BORDER
        ws.column_dimensions[letra].width = ancho

    fila = 2
    for f in facturas:
        # Usar el número de línea guardado, o el índice
        num_linea = f.get("linea", 0) or (fila - 1)
        ws[f"A{fila}"] = num_linea
        ws[f"A{fila}"].alignment = CENTRO
        ws[f"A{fila}"].border = BORDER

        ws[f"B{fila}"] = f.get("fecha", "")
        ws[f"B{fila}"].alignment = CENTRO
        ws[f"B{fila}"].border = BORDER

        # FACTURA (con enlace al PDF)
        c = ws[f"C{fila}"]
        c.value = f.get("folio", "")
        c.alignment = IZQUIERDA
        c.border = BORDER
        ruta_pdf = _obtener_ruta_pdf_egreso(f, num_linea)
        if ruta_pdf:
            c.hyperlink = ruta_pdf
            c.font = FUENTE_LINK

        # QVET (vacío, manual)
        ws[f"D{fila}"] = ""
        ws[f"D{fila}"].alignment = CENTRO
        ws[f"D{fila}"].border = BORDER

        # PROVEEDOR
        ws[f"E{fila}"] = f.get("nombre_emisor", "")
        ws[f"E{fila}"].alignment = IZQUIERDA
        ws[f"E{fila}"].border = BORDER

        # UUID
        ws[f"F{fila}"] = f.get("uuid", "")
        ws[f"F{fila}"].alignment = IZQUIERDA
        ws[f"F{fila}"].border = BORDER

        # METODO
        ws[f"G{fila}"] = "PPD"
        ws[f"G{fila}"].alignment = CENTRO
        ws[f"G{fila}"].border = BORDER

        # SUBTOTAL
        ws[f"H{fila}"] = float(f.get("subtotal", 0))
        ws[f"H{fila}"].number_format = FORMATO_MONEDA
        ws[f"H{fila}"].alignment = DERECHA
        ws[f"H{fila}"].border = BORDER

        # IVA
        ws[f"I{fila}"] = float(f.get("iva", 0))
        ws[f"I{fila}"].number_format = FORMATO_MONEDA
        ws[f"I{fila}"].alignment = DERECHA
        ws[f"I{fila}"].border = BORDER

        # IEPS
        ws[f"J{fila}"] = float(f.get("ieps", 0))
        ws[f"J{fila}"].number_format = FORMATO_MONEDA
        ws[f"J{fila}"].alignment = DERECHA
        ws[f"J{fila}"].border = BORDER

        # TOTAL
        ws[f"K{fila}"] = float(f.get("total", 0))
        ws[f"K{fila}"].number_format = FORMATO_MONEDA
        ws[f"K{fila}"].alignment = DERECHA
        ws[f"K{fila}"].border = BORDER

        # REFERENCIA (vacío, manual)
        ws[f"L{fila}"] = ""
        ws[f"L{fila}"].alignment = CENTRO
        ws[f"L{fila}"].border = BORDER

        fila += 1

    # Totales
    ws[f"E{fila}"] = "TOTAL:"
    ws[f"E{fila}"].font = BOLD_DARK
    ws[f"E{fila}"].alignment = DERECHA
    for letra in ["H", "I", "J", "K"]:
        c = ws[f"{letra}{fila}"]
        c.value = f"=SUM({letra}2:{letra}{fila-1})"
        c.font = BOLD_DARK
        c.number_format = FORMATO_MONEDA
        c.alignment = DERECHA

    wb.save(ruta_xlsx)
    return ruta_xlsx


def _obtener_ruta_pdf_egreso(factura, num_linea):
    """
    Construye la ruta relativa al PDF de una factura de Egreso.
    El Excel está en Egreso/Deposito_Egreso/.
    El PDF está en Egreso/PUE/ o Egreso/PPD/.
    """
    try:
        folio = factura.get("folio", "")
        if not folio:
            return None

        # Determinar subcarpeta
        metodo = factura.get("metodo_pago", "PUE")
        if metodo == "PPD":
            subcarpeta = "PPD"
        else:
            subcarpeta = "PUE"

        # Ruta relativa desde Deposito_Egreso a PUE/ o PPD/
        return f"../{subcarpeta}/{num_linea}-{folio}.PDF"
    except Exception:
        return None
