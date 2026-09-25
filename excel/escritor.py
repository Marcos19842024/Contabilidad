# -*- coding: utf-8 -*-
"""
excel/escritor.py
Escritura de filas, encabezados y totales en una hoja de Excel.
"""

from datetime import datetime

from openpyxl.styles import PatternFill

from excel.constantes import (
    COLORES_SECCION_EXCEL,
    COLUMNAS_EXCEL,
    MAPA_CLAVES_EXCEL,
    FORMULAS_AUTO_EXCEL,
    GRUPOS_EXCEL,
)
from excel.estilos import (
    BORDER,
    CENTRO,
    DERECHA,
    IZQUIERDA,
    BOLD_WHITE,
    BOLD_DARK,
    FUENTE_NORMAL,
    SIN_RELLENO,
    FILL_TOTAL,
    FORMATO_MONEDA,
    FORMATO_FECHA,
)


def escribir_encabezados(ws):
    """Escribe fila 1 (grupos) y fila 2 (columnas) con estilos."""
    # --- Fila 1: grupos ---
    for ini, fin, texto in GRUPOS_EXCEL:
        c = ws.cell(row=1, column=ini, value=texto)
        c.font = BOLD_DARK
        c.alignment = CENTRO

        for col in range(ini, fin + 1):
            celda = ws.cell(row=1, column=col)
            celda.border = BORDER
            celda.alignment = CENTRO

        if ini != fin:
            ws.merge_cells(start_row=1, start_column=ini,
                           end_row=1, end_column=fin)

    # --- Fila 2: columnas ---
    for letra, (titulo, seccion, tipo) in COLUMNAS_EXCEL.items():
        c = ws[f"{letra}2"]
        c.value = titulo
        c.font = BOLD_WHITE
        c.fill = PatternFill("solid", fgColor=COLORES_SECCION_EXCEL[seccion])
        c.alignment = CENTRO
        c.border = BORDER


def _buscar_letra_por_clave(clave_buscada):
    """Devuelve la letra de columna cuya clave coincide."""
    for letra, clave in MAPA_CLAVES_EXCEL.items():
        if clave == clave_buscada:
            return letra
    return None


def escribir_fila(ws, r, fila):
    """Escribe UNA fila de registro con estilos, fórmulas y casos especiales."""
    valor_tc = float(r.get("tc", 0) or 0)
    valor_td = float(r.get("td", 0) or 0)
    valor_tarjeta = valor_tc + valor_td
    valor_efectivo = float(r.get("efectivo", 0) or 0)
    valor_transfer = float(r.get("transfer", 0) or 0)

    for letra, (titulo, seccion, tipo) in COLUMNAS_EXCEL.items():
        clave = MAPA_CLAVES_EXCEL[letra]
        c = ws[f"{letra}{fila}"]

        # Reset de estilos base
        c.font = FUENTE_NORMAL
        c.fill = SIN_RELLENO
        c.border = BORDER

        # --- Casos especiales ---
        if clave == "__tarjeta__":
            letra_tc = _buscar_letra_por_clave("tc")
            letra_td = _buscar_letra_por_clave("td")
            if letra_tc and letra_td:
                c.value = f"={letra_tc}{fila}+{letra_td}{fila}"
                c.number_format = FORMATO_MONEDA
            elif valor_tarjeta > 0:
                c.value = valor_tarjeta
                c.number_format = FORMATO_MONEDA
            c.alignment = DERECHA
            continue

        if clave == "__efectivo_dup__":
            letra_efectivo = _buscar_letra_por_clave("efectivo")
            if letra_efectivo:
                c.value = f"={letra_efectivo}{fila}"
                c.number_format = FORMATO_MONEDA
            elif valor_efectivo > 0:
                c.value = valor_efectivo
                c.number_format = FORMATO_MONEDA
            c.alignment = DERECHA
            continue

        if clave == "__transfer_dup__":
            letra_transfer = _buscar_letra_por_clave("transfer")
            if letra_transfer:
                c.value = f"={letra_transfer}{fila}"
                c.number_format = FORMATO_MONEDA
            elif valor_transfer > 0:
                c.value = valor_transfer
                c.number_format = FORMATO_MONEDA
            c.alignment = DERECHA
            continue

        if clave == "":
            c.alignment = CENTRO
            continue

        # --- Casos normales ---
        valor = r.get(clave, 0 if tipo == "money" else "")
        expresion = r.get(f"{clave}__expr", "")

        if tipo == "money":
            if expresion:
                c.value = f"={expresion}"
            elif letra in FORMULAS_AUTO_EXCEL:
                c.value = FORMULAS_AUTO_EXCEL[letra].format(r=fila)
            else:
                c.value = float(valor or 0)
            c.number_format = FORMATO_MONEDA
            c.alignment = DERECHA

        elif tipo == "date":
            if valor:
                try:
                    fecha_dt = datetime.strptime(str(valor).strip(), "%d/%m/%Y")
                    c.value = fecha_dt
                    c.number_format = FORMATO_FECHA
                except (ValueError, TypeError):
                    c.value = valor
            c.alignment = CENTRO

        else:
            c.value = valor
            c.alignment = IZQUIERDA


def escribir_fila_totales(ws, fila):
    """Escribe la fila TOTALES con fórmulas =SUM(...)."""
    MERGE_INI = 1
    MERGE_FIN = 5

    # Etiqueta "TOTALES" en la primera celda
    c_tot = ws.cell(row=fila, column=MERGE_INI, value="TOTALES")
    c_tot.font = BOLD_DARK
    c_tot.fill = FILL_TOTAL
    c_tot.alignment = CENTRO
    c_tot.border = BORDER

    # Estilos del rango combinado
    for col in range(MERGE_INI, MERGE_FIN + 1):
        celda = ws.cell(row=fila, column=col)
        try:
            celda.fill = FILL_TOTAL
            celda.border = BORDER
            celda.alignment = CENTRO
        except Exception:
            pass

    ws.merge_cells(start_row=fila, start_column=MERGE_INI,
                   end_row=fila, end_column=MERGE_FIN)

    # Fórmulas SUM en el resto
    for letra, (titulo, seccion, tipo) in COLUMNAS_EXCEL.items():
        col_idx = ws[f"{letra}1"].column
        if MERGE_INI <= col_idx <= MERGE_FIN:
            continue
        c = ws.cell(row=fila, column=col_idx)
        if tipo == "money":
            c.value = f"=SUM({letra}3:{letra}{fila - 1})"
            c.number_format = FORMATO_MONEDA
            c.alignment = DERECHA
        c.font = BOLD_DARK
        c.fill = FILL_TOTAL
        c.border = BORDER


def ajustar_anchos(ws, fila_total):
    """Ajusta el ancho de las columnas según el contenido."""
    from openpyxl.utils import get_column_letter

    for letra in COLUMNAS_EXCEL.keys():
        largo_max = len(str(COLUMNAS_EXCEL[letra][0]))
        for ini, fin, texto in GRUPOS_EXCEL:
            col_ini = get_column_letter(ini)
            col_fin = get_column_letter(fin)
            if col_ini <= letra <= col_fin:
                n = fin - ini + 1
                largo_max = max(largo_max, len(texto) / n if n else 0)
                break
        for f in range(3, fila_total):
            celda = ws[f"{letra}{f}"]
            if celda.value is None:
                continue
            if isinstance(celda.value, datetime):
                txt = celda.value.strftime("%d/%m/%Y")
            elif isinstance(celda.value, (int, float)):
                txt = f"${celda.value:,.2f}"
            else:
                txt = str(celda.value)
            largo_max = max(largo_max, len(txt))
        ancho = min(max(largo_max + 2, 8), 40)
        ws.column_dimensions[letra].width = ancho

    ws.column_dimensions["AP"].width = max(ws.column_dimensions["AP"].width, 38)
    ws.column_dimensions["D"].width = max(ws.column_dimensions["D"].width, 22)