# -*- coding: utf-8 -*-
"""
excel/estilos.py
Estilos comunes para las celdas del Excel.
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# --- Bordes ---
THIN = Side(border_style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# --- Alineaciones ---
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
DERECHA = Alignment(horizontal="right", vertical="center")
IZQUIERDA = Alignment(horizontal="left", vertical="center")

# --- Fuentes ---
BOLD_WHITE = Font(bold=True, color="FFFFFF", size=11)
BOLD_DARK = Font(bold=True, color="000000", size=10)
FUENTE_NORMAL = Font()

# --- Rellenos ---
SIN_RELLENO = PatternFill()
FILL_TOTAL = PatternFill("solid", fgColor="FFD966")

# --- Formatos de número ---
FORMATO_MONEDA = '"$"#,##0.00'
FORMATO_FECHA = "DD/MM/YYYY"