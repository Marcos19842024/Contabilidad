# -*- coding: utf-8 -*-
"""
excel/generador.py
Orquestación: crear Excel nuevo, anexar registros a uno existente,
y detección de archivo bloqueado.
"""

import sys
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from config.campos import MESES_ES
from excel.constantes import COLUMNAS_EXCEL
from excel.escritor import (
    escribir_encabezados,
    escribir_fila,
    escribir_fila_totales,
    ajustar_anchos,
)


# ============================================================
# DETECCIÓN DE ARCHIVO BLOQUEADO
# ============================================================
def archivo_esta_bloqueado(ruta):
    """
    Devuelve True si el archivo existe y está bloqueado por otro programa.
    """
    import os
    ruta = Path(ruta)
    if not ruta.exists():
        return False

    if sys.platform.startswith("win"):
        try:
            os.rename(str(ruta), str(ruta))
            return False
        except OSError:
            return True
    else:
        try:
            with open(ruta, "r+b"):
                pass
            return False
        except (OSError, IOError):
            return True


# ============================================================
# CREAR EXCEL NUEVO
# ============================================================
def escribir_excel(ruta, regs, var_mes=None, var_anio=None):
    """
    Crea un Excel nuevo (sobrescribe si existe).
    `var_mes` y `var_anio` son opcionales (para el nombre de la hoja si regs está vacío).
    """
    if regs:
        mes_actual = regs[0].get("mes", "") or var_mes or MESES_ES[datetime.now().month - 1]
        anio_actual = regs[0].get("anio", "") or var_anio or datetime.now().year
    else:
        mes_actual = var_mes or MESES_ES[datetime.now().month - 1]
        anio_actual = var_anio or datetime.now().year

    nombre_hoja = f"{str(mes_actual).capitalize()} {anio_actual}"[:31]

    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja

    escribir_encabezados(ws)

    fila = 3
    for r in regs:
        escribir_fila(ws, r, fila)
        fila += 1

    escribir_fila_totales(ws, fila)
    ajustar_anchos(ws, fila)

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "F3"

    wb.save(ruta)


# ============================================================
# ANEXAR A EXCEL EXISTENTE
# ============================================================
def anexar_al_excel(ruta, regs):
    """
    Abre un Excel existente y anexa solo los registros nuevos.
    Detecta duplicados por:
      1. Folio fiscal (UUID) — columna AP
      2. No. de factura      — columna A
    Respeta ediciones manuales de las filas existentes.
    Devuelve (nuevos_agregados, omitidos).
    """
    wb = load_workbook(ruta)
    ws = wb.active

    COL_FOLIO_FISCAL = 42  # AP
    COL_NO_FACTURA = 1     # A

    # --- 1. Leer lo existente ---
    folios_existentes = set()
    no_facturas_existentes = set()
    fila_totales_original = None

    fila = 3
    max_fila = ws.max_row + 10
    while fila <= max_fila:
        val_no_fac = ws.cell(row=fila, column=COL_NO_FACTURA).value
        val_folio = ws.cell(row=fila, column=COL_FOLIO_FISCAL).value

        if val_no_fac is not None and str(val_no_fac).strip().upper() == "TOTALES":
            fila_totales_original = fila
            break

        if val_folio:
            folios_existentes.add(str(val_folio).strip().upper())
        if val_no_fac:
            no_facturas_existentes.add(str(val_no_fac).strip().upper())
        fila += 1

    if fila_totales_original is None:
        fila_totales_original = fila

    # --- 2. Filtrar nuevos ---
    nuevos = []
    omitidos = 0
    for r in regs:
        uuid = str(r.get("folio_fiscal", "")).strip().upper()
        no_fac = str(r.get("no_factura", "")).strip().upper()

        es_duplicado = False
        if uuid and uuid in folios_existentes:
            es_duplicado = True
        elif no_fac and no_fac in no_facturas_existentes:
            es_duplicado = True

        if es_duplicado:
            omitidos += 1
            continue
        nuevos.append(r)

    if not nuevos:
        return 0, omitidos

    # --- 3. Deshacer merges de la fila TOTALES ---
    merges_a_quitar = []
    for rango in list(ws.merged_cells.ranges):
        if rango.min_row == fila_totales_original:
            merges_a_quitar.append(str(rango))
    for rango_str in merges_a_quitar:
        try:
            ws.unmerge_cells(rango_str)
        except Exception:
            pass

    # --- 4. Respaldar la fila TOTALES ---
    totales_guardados = {}
    if fila_totales_original is not None:
        for col in range(1, ws.max_column + 1):
            celda = ws.cell(row=fila_totales_original, column=col)
            totales_guardados[col] = {
                "value": celda.value,
                "font": copy(celda.font),
                "fill": copy(celda.fill),
                "border": copy(celda.border),
                "alignment": copy(celda.alignment),
                "number_format": celda.number_format,
            }
        # Limpiar valores Y estilos
        from openpyxl.styles import Font, PatternFill, Alignment, Border
        for col in range(1, ws.max_column + 1):
            celda = ws.cell(row=fila_totales_original, column=col)
            celda.value = None
            celda.font = Font()
            celda.fill = PatternFill()
            celda.border = Border()
            celda.alignment = Alignment()
            celda.number_format = "General"

    # --- 5. Escribir nuevas filas ---
    fila_insercion = fila_totales_original
    for r in nuevos:
        escribir_fila(ws, r, fila_insercion)
        fila_insercion += 1

    # --- 6. Reescribir TOTALES ---
    fila_totales_nueva = fila_insercion
    escribir_fila_totales(ws, fila_totales_nueva)

    # --- 7. Ajustar anchos ---
    ajustar_anchos(ws, fila_totales_nueva)

    wb.save(ruta)
    return len(nuevos), omitidos