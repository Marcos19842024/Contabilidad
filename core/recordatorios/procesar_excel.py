# core/recordatorios/procesar_excel.py
"""
Procesamiento de los Excel de agenda y vacunas.
"""

from openpyxl import load_workbook

from core.recordatorios.formatear import (
    format_string,
    format_numbers,
    normalizar_telefono,
    format_date_long,
    format_fecha_para_mensaje,
    extraer_nombre_clinica,
)
from core.recordatorios.servicios import procesar_servicio


# ---------------------------------------------------------------
# Detección de tipo de Excel
# ---------------------------------------------------------------

def _quitar_acentos(texto: str) -> str:
    if not texto:
        return ""
    reemplazos = {
        "Á": "A", "É": "E", "Í": "I", "Ó": "O", "Ú": "U",
        "À": "A", "È": "E", "Ì": "I", "Ò": "O", "Ù": "U",
        "Ñ": "N", "Ü": "U",
    }
    for a, b in reemplazos.items():
        texto = texto.replace(a, b)
    return texto


def _leer_encabezados_normalizados(ws) -> list:
    encabezados = []
    for cell in ws[1]:
        val = cell.value
        if val is None:
            encabezados.append("")
        else:
            encabezados.append(_quitar_acentos(str(val).strip().upper()))
    return encabezados


def _tiene_columnas(encabezados: list, requeridas: list) -> bool:
    for req in requeridas:
        req_norm = _quitar_acentos(req.upper())
        if req_norm not in encabezados:
            return False
    return True


def _indice_de(encabezados: list, nombre: str) -> int:
    nombre_norm = _quitar_acentos(nombre.upper())
    for i, h in enumerate(encabezados):
        if h == nombre_norm:
            return i
    return -1


def detectar_tipo_excel(ruta_excel: str):
    try:
        wb = load_workbook(ruta_excel, data_only=True, read_only=True)
        ws = wb.active
        encabezados = _leer_encabezados_normalizados(ws)
        wb.close()

        if _tiene_columnas(encabezados, ["TIPO VISITA", "INICIO", "PROPIETARIO"]):
            return "citas"
        if _tiene_columnas(encabezados, ["TIPO DE RECORDATORIO", "VACUNA", "PRÓXIMA FECHA"]):
            return "vacunas"
    except Exception as e:
        print(f"[detectar_tipo_excel] Error: {e}")
    return None


# ---------------------------------------------------------------
# Procesamiento de citas
# ---------------------------------------------------------------

def _procesar_citas(rows: list, nombre_clinica: str, indices: dict) -> list:
    def _get(fila, col):
        i = indices.get(col, -1)
        if i < 0 or i >= len(fila):
            return ""
        return fila[i]

    citas_por_cliente = {}

    for idx, fila in enumerate(rows):
        try:
            fecha_raw = _get(fila, "FECHA")
            hora = str(_get(fila, "INICIO") or "").strip()
            tipo_visita_raw = str(_get(fila, "TIPO VISITA") or "").strip()
            propietario_raw = str(_get(fila, "PROPIETARIO") or "").strip()
            mascota_raw = str(_get(fila, "MASCOTA") or "").strip()
            telefono_raw = str(_get(fila, "TELÉFONO") or "").strip()
            asunto = str(_get(fila, "ASUNTO") or "").strip()
            agenda = str(_get(fila, "AGENDA") or "").strip()
            estado = str(_get(fila, "ESTADO") or "").strip()

            propietario = format_string(propietario_raw)
            telefono_display = telefono_raw
            telefono = normalizar_telefono(telefono_raw)
            fecha = format_date_long(fecha_raw)

            if not propietario or not telefono:
                continue

            clave = f"{propietario}_{telefono}"
            if clave not in citas_por_cliente:
                citas_por_cliente[clave] = {
                    "nombre": propietario,
                    "telefono": telefono,
                    "telefono_display": telefono_display,
                    "citas": [],
                }

            citas_por_cliente[clave]["citas"].append({
                "fecha": fecha,
                "hora": hora,
                "tipo_visita": format_string(tipo_visita_raw),
                "tipo_visita_raw": tipo_visita_raw,
                "mascota": format_string(mascota_raw),
                "asunto": asunto,
                "agenda": agenda,
                "estado": estado,
            })
        except Exception as e:
            print(f"Error procesando fila {idx + 2}: {e}")
            continue

    clientes = []
    for clave, data in citas_por_cliente.items():
        cliente = {
            "nombre": data["nombre"],
            "telefono": data["telefono"],
            "telefono_display": data["telefono_display"],
            "citas": data["citas"],
            "mascotas": list({c["mascota"] for c in data["citas"] if c["mascota"]}),
            "mensajes": [],
            "tipo": "citas",
        }
        cliente["mensajes"] = _generar_mensajes_citas(cliente, nombre_clinica)
        clientes.append(cliente)

    return clientes


def _generar_mensajes_citas(cliente: dict, nombre_clinica: str) -> list:
    citas = cliente["citas"]
    if not citas:
        return []

    nombre_clinica_limpio = extraer_nombre_clinica(nombre_clinica)
    mensajes = []

    mensajes.append(f"Hola {cliente['nombre']}.")

    primera_fecha = citas[0]["fecha"]
    mismo_dia = all(c["fecha"] == primera_fecha for c in citas)

    citas_por_mascota = {}
    for c in citas:
        if c["mascota"]:
            citas_por_mascota.setdefault(c["mascota"], []).append(c)

    mascotas_con_citas = list(citas_por_mascota.keys())

    if not mascotas_con_citas:
        mensajes.append(f"{nombre_clinica_limpio} le recuerda su cita.")
        return mensajes

    incluir_fecha = not mismo_dia

    if len(mascotas_con_citas) == 1:
        mascota_nombre = mascotas_con_citas[0]
        citas_mascota = citas_por_mascota[mascota_nombre]
        mensaje = f"la cita de su mascota '{mascota_nombre}' "
        mensaje += _listar_citas_para_mascota(citas_mascota, incluir_fecha)
    else:
        mensaje = "las citas de sus mascotas: "
        for i, mascota_nombre in enumerate(mascotas_con_citas):
            citas_mascota = citas_por_mascota[mascota_nombre]
            if i == 0:
                mensaje += f"'{mascota_nombre}' "
                mensaje += _listar_citas_para_mascota(citas_mascota, incluir_fecha)
            elif i == len(mascotas_con_citas) - 1:
                mensaje += f" y '{mascota_nombre}' "
                mensaje += _listar_citas_para_mascota(citas_mascota, incluir_fecha)
            else:
                mensaje += f", '{mascota_nombre}' "
                mensaje += _listar_citas_para_mascota(citas_mascota, incluir_fecha)

    mensaje_completo = f"{nombre_clinica_limpio} le recuerda {mensaje}"

    if mismo_dia and primera_fecha:
        fecha_formateada = format_fecha_para_mensaje(primera_fecha)
        if fecha_formateada == "mañana":
            mensaje_completo += " mañana"
        else:
            mensaje_completo += f" {fecha_formateada}"

    mensaje_completo += ".\n\n"

    if len(citas) > 1:
        horas_unicas = sorted({c["hora"] for c in citas if c["hora"]})
        if horas_unicas:
            if len(horas_unicas) == 1:
                mensaje_completo += f"⏰ Hora: {horas_unicas[0]}\n"
            else:
                mensaje_completo += f"⏰ Horas: {', '.join(horas_unicas)}\n"
    elif citas[0]["hora"]:
        mensaje_completo += f"⏰ Hora: {citas[0]['hora']}\n"

    mensaje_completo += "\nPor favor confirme su asistencia y el servicio con anticipación.\n\n¡Gracias! 🐾"

    mensajes.append(mensaje_completo)
    return mensajes


def _listar_citas_para_mascota(citas: list, incluir_fecha: bool) -> str:
    if not citas:
        return ""
    citas_ordenadas = sorted(citas, key=lambda c: c.get("hora", "") or "99:99")

    if len(citas_ordenadas) == 1:
        return _describir_cita(citas_ordenadas[0], incluir_fecha)

    descripcion = f"tiene {len(citas_ordenadas)} citas programadas: "
    for i, cita in enumerate(citas_ordenadas):
        servicio = procesar_servicio(
            cita["tipo_visita_raw"], cita["asunto"], cita["estado"]
        )
        if i == 0:
            descripcion += _construir_descripcion_cita(cita, servicio, incluir_fecha, False)
        elif i == len(citas_ordenadas) - 1:
            descripcion += f" y {_construir_descripcion_cita(cita, servicio, incluir_fecha, True)}"
        else:
            descripcion += f", {_construir_descripcion_cita(cita, servicio, incluir_fecha, True)}"

    return descripcion


def _describir_cita(cita: dict, incluir_fecha: bool) -> str:
    servicio = procesar_servicio(cita["tipo_visita_raw"], cita["asunto"], cita["estado"])
    return _construir_descripcion_cita(cita, servicio, incluir_fecha, False)


def _construir_descripcion_cita(cita: dict, servicio: str, incluir_fecha: bool, es_segunda_o_mas: bool) -> str:
    descripcion = ""
    tipo_visita_base = (cita["tipo_visita_raw"] or "").lower()
    es_estetica = "peluquer" in tipo_visita_base or "estética" in tipo_visita_base or "estetica" in tipo_visita_base

    if es_estetica:
        if es_segunda_o_mas:
            descripcion += servicio
        else:
            descripcion += f"para {servicio}"
    else:
        tipo_formateado = cita["tipo_visita"]
        if tipo_visita_base and tipo_visita_base in servicio.lower():
            if es_segunda_o_mas:
                descripcion += servicio
            else:
                descripcion += f"para {servicio}"
        else:
            if servicio:
                if es_segunda_o_mas:
                    descripcion += f"{tipo_formateado} ({servicio})"
                else:
                    descripcion += f"para {tipo_formateado} ({servicio})"
            else:
                if es_segunda_o_mas:
                    descripcion += tipo_formateado
                else:
                    descripcion += f"para {tipo_formateado}"

    if cita["hora"]:
        descripcion += f" a las {cita['hora']}"

    if incluir_fecha and cita["fecha"]:
        descripcion += f" {format_fecha_para_mensaje(cita['fecha'])}"

    return descripcion


# ---------------------------------------------------------------
# Procesamiento de vacunas
# ---------------------------------------------------------------

def _procesar_vacunas(rows: list, nombre_clinica: str, indices: dict) -> list:
    def _get(fila, col):
        i = indices.get(col, -1)
        if i < 0 or i >= len(fila):
            return ""
        return fila[i]

    clientes_map = {}

    for idx, fila in enumerate(rows):
        try:
            nombre_cliente_raw = str(_get(fila, "CLIENTE") or "").strip()
            telefono_raw = str(_get(fila, "TELÉFONO 1") or "").strip()
            mascota_raw = str(_get(fila, "MASCOTA") or "").strip()
            recordatorio_raw = str(_get(fila, "TIPO DE RECORDATORIO") or "").strip()
            tipo_raw = str(_get(fila, "VACUNA") or "").strip()
            fecha_raw = _get(fila, "PRÓXIMA FECHA")

            nombre_cliente = format_string(nombre_cliente_raw)
            telefono = normalizar_telefono(telefono_raw)

            if not nombre_cliente or not telefono:
                continue

            fecha = ""
            if fecha_raw:
                if hasattr(fecha_raw, "strftime"):
                    fecha = fecha_raw.strftime("%Y-%m-%d")
                else:
                    fecha = str(fecha_raw).strip()

            clave = f"{nombre_cliente}_{telefono}"
            if clave not in clientes_map:
                clientes_map[clave] = {
                    "nombre": nombre_cliente,
                    "telefono": telefono,
                    "telefono_display": telefono_raw,
                    "mascotas": {},
                }

            mascota_nombre = format_string(mascota_raw)
            if not mascota_nombre:
                continue

            if mascota_nombre not in clientes_map[clave]["mascotas"]:
                clientes_map[clave]["mascotas"][mascota_nombre] = {}

            recordatorio_nombre = format_string(recordatorio_raw)
            if not recordatorio_nombre:
                continue

            if recordatorio_nombre not in clientes_map[clave]["mascotas"][mascota_nombre]:
                clientes_map[clave]["mascotas"][mascota_nombre][recordatorio_nombre] = []

            tipo_nombre = format_string(tipo_raw)
            if tipo_nombre and fecha:
                tipos = clientes_map[clave]["mascotas"][mascota_nombre][recordatorio_nombre]
                if not any(t["nombre"] == tipo_nombre and t["fecha"] == fecha for t in tipos):
                    tipos.append({"nombre": tipo_nombre, "fecha": fecha})

        except Exception as e:
            print(f"Error procesando fila {idx + 2}: {e}")
            continue

    clientes = []
    for clave, data in clientes_map.items():
        cliente = {
            "nombre": data["nombre"],
            "telefono": data["telefono"],
            "telefono_display": data["telefono_display"],
            "mascotas": [],
            "mensajes": [],
            "tipo": "vacunas",
        }
        for nombre_mascota, recordatorios_dict in data["mascotas"].items():
            mascota = {
                "nombre": nombre_mascota,
                "recordatorios": [],
            }
            for nombre_record, tipos in recordatorios_dict.items():
                mascota["recordatorios"].append({
                    "nombre": nombre_record,
                    "tipos": tipos,
                })
            cliente["mascotas"].append(mascota)

        cliente["mensajes"] = _generar_mensajes_vacunas(cliente, nombre_clinica)
        clientes.append(cliente)

    return clientes


def _generar_mensajes_vacunas(cliente: dict, nombre_clinica: str) -> list:
    mascotas = cliente["mascotas"]
    mensajes = []

    mensajes.append(f"Hola {cliente['nombre']}.")

    if not mascotas:
        return mensajes

    if len(mascotas) == 1:
        mensaje = f"su mascota '{mascotas[0]['nombre']}'" + _listar_recordatorios(mascotas[0])
    else:
        mensaje = "sus mascotas: "
        for i, m in enumerate(mascotas):
            if i == 0:
                mensaje += f"'{m['nombre']}'" + _listar_recordatorios(m)
            elif i == len(mascotas) - 1:
                mensaje += f" y '{m['nombre']}'" + _listar_recordatorios(m)
            else:
                mensaje += f", '{m['nombre']}'" + _listar_recordatorios(m)

    if (mascotas and mascotas[0]["recordatorios"] and
            mascotas[0]["recordatorios"][0]["tipos"]):
        fecha = mascotas[0]["recordatorios"][0]["tipos"][0]["fecha"]
        if fecha:
            mensaje += f" el día {format_date_long(fecha)}."
        else:
            mensaje += "."
    else:
        mensaje += "."

    nombre_clinica_limpio = extraer_nombre_clinica(nombre_clinica)
    mensajes.append(f"{nombre_clinica_limpio} le informa que {mensaje}")
    mensajes.append("\n🐾 ¿Quiere agendar su cita?")

    return mensajes


def _listar_recordatorios(mascota: dict) -> str:
    texto = " tiene pendiente "
    recordatorios = mascota["recordatorios"]

    if not recordatorios:
        return texto.strip()

    if len(recordatorios) == 1:
        texto += recordatorios[0]["nombre"] + _listar_tipos(recordatorios[0])
    else:
        for i, r in enumerate(recordatorios):
            if i == 0:
                texto += r["nombre"] + _listar_tipos(r)
            elif i == len(recordatorios) - 1:
                texto += " y " + r["nombre"] + _listar_tipos(r)
            else:
                texto += ", " + r["nombre"] + _listar_tipos(r)

    return texto


def _listar_tipos(recordatorio: dict) -> str:
    tipos = recordatorio["tipos"]
    if not tipos:
        return ""

    if len(tipos) == 1:
        return f" ({tipos[0]['nombre']})"

    texto = " ("
    for i, t in enumerate(tipos):
        if i == 0:
            texto += t["nombre"]
        elif i == len(tipos) - 1:
            texto += f" y {t['nombre']}"
        else:
            texto += f", {t['nombre']}"
    texto += ")"
    return texto


# ---------------------------------------------------------------
# Función principal
# ---------------------------------------------------------------

def procesar_excel(ruta_excel: str, nombre_clinica: str) -> dict:
    tipo = detectar_tipo_excel(ruta_excel)
    if not tipo:
        raise ValueError(
            "El Excel no coincide con ningún formato conocido.\n\n"
            "Formatos soportados:\n"
            "• Citas: FECHA | INICIO | TIPO VISITA | PROPIETARIO | MASCOTA | "
            "TELÉFONO | ASUNTO | AGENDA | ESTADO\n"
            "• Vacunas: CLIENTE | TELÉFONO 1 | MASCOTA | TIPO DE RECORDATORIO | "
            "VACUNA | PRÓXIMA FECHA"
        )

    wb = load_workbook(ruta_excel, data_only=True)
    ws = wb.active

    encabezados = _leer_encabezados_normalizados(ws)

    indices = {}
    if tipo == "citas":
        for col in ["FECHA", "INICIO", "TIPO VISITA", "PROPIETARIO", "MASCOTA",
                    "TELÉFONO", "ASUNTO", "AGENDA", "ESTADO"]:
            indices[col] = _indice_de(encabezados, col)
    else:
        for col in ["CLIENTE", "TELÉFONO 1", "MASCOTA",
                    "TIPO DE RECORDATORIO", "VACUNA", "PRÓXIMA FECHA"]:
            indices[col] = _indice_de(encabezados, col)

    filas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None:
            continue
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        filas.append(row)
    wb.close()

    if not filas:
        raise ValueError("El Excel no contiene datos")

    if tipo == "citas":
        clientes = _procesar_citas(filas, nombre_clinica, indices)
    else:
        clientes = _procesar_vacunas(filas, nombre_clinica, indices)

    return {
        "tipo": tipo,
        "clientes": clientes,
        "total": len(clientes),
    }


# ---------------------------------------------------------------
# Regenerar mensajes (al cambiar sucursal)
# ---------------------------------------------------------------

def regenerar_mensajes_cliente(cliente: dict, nombre_clinica: str):
    """
    Regenera los mensajes de un cliente existente con un nuevo nombre de clínica.
    Se usa cuando el usuario cambia la sucursal activa.
    """
    if cliente.get("tipo") == "citas":
        cliente["mensajes"] = _generar_mensajes_citas(cliente, nombre_clinica)
    elif cliente.get("tipo") == "vacunas":
        cliente["mensajes"] = _generar_mensajes_vacunas(cliente, nombre_clinica)